"""Market business report export (Q-M6  2026-10-11).

已冻结出清结果的业务投影：Excel＋HTML，与成本/结算报告同款纪律（内存
构建、不可变文件、清单最后写、同身份重放按哈希比对）。报告只投影冻结
记录——不重算出清、不读台账最新版本。
"""

from __future__ import annotations

import hashlib
import io
import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

_MANIFEST_FILE = "market_report_manifest.json"
_REPORT_XLSX = "market_report.xlsx"
_REPORT_HTML = "market_report.html"
_GENERATOR = "market-report-1.0"


class MarketReportError(ValueError):
    """A market-report export refusal (code + message)."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _canonical(payload: Any) -> bytes:
    return json.dumps(
        payload, ensure_ascii=False, sort_keys=True, default=str
    ).encode("utf-8")


def _fmt(value: Any, digits: int = 2) -> str:
    if isinstance(value, bool) or value is None:
        return str(value)
    if isinstance(value, (int, float)):
        return f"{value:,.{digits}f}"
    return str(value)


_STATUS_LABELS = {
    "cleared": "出清成功",
    "cleared_with_shortfall": "触顶出清（有缺口）",
    "not_cleared": "未出清（无有效供给）",
}


def _excel_bytes(*, formal: dict[str, Any]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    # ① 对象与结论
    ws = wb.active
    ws.title = "对象与结论"
    ws["B2"] = "容量市场仿真出清报告（正式结果业务投影）"
    ws["B2"].font = bold
    rows = [
        ("出清编号", formal.get("clearing_ref")),
        ("规则版本", formal.get("rule_version")),
        ("确认人", formal.get("confirmer")),
        ("冻结时间", formal.get("frozen_at")),
        ("出清状态", _STATUS_LABELS.get(str(formal.get("status")), formal.get("status"))),
        ("统一出清价（元/kW·年）", formal.get("clearing_price")),
        ("中标容量（MW）", formal.get("procured_mw")),
        ("目标容量（MW）", formal.get("target_mw")),
        ("最低采购要求（MW）", formal.get("min_procurement_mw")),
        ("容量缺口（MW）", formal.get("shortfall_mw")),
        ("采购完成率", formal.get("completion_ratio")),
        ("并列报价按比例分配", "是" if formal.get("tie_allocated") else "否"),
    ]
    for i, (label, value) in enumerate(rows, start=4):
        ws.cell(row=i, column=2, value=label).font = bold
        ws.cell(
            row=i,
            column=3,
            value=(
                _fmt(value)
                if isinstance(value, (int, float)) and not isinstance(value, bool)
                else str(value)
            ),
        )

    # ② 需求曲线快照
    curve = formal.get("curve") or {}
    ws2 = wb.create_sheet("需求曲线")
    ws2["B2"] = "VRR 三点需求曲线（锚＝目标容量×Net CONE）"
    ws2["B2"].font = bold
    headers = ["容量（MW）", "价格（元/kW·年）"]
    for col, header in enumerate(headers, start=2):
        ws2.cell(row=4, column=col, value=header).font = bold
    for i, (q, p) in enumerate(curve.get("points") or [], start=5):
        ws2.cell(row=i, column=2, value=_fmt(q))
        ws2.cell(row=i, column=3, value=_fmt(p))
    ws2.cell(row=9, column=2, value="Net CONE（元/kW·年）").font = bold
    ws2.cell(row=9, column=3, value=_fmt(curve.get("net_cone")))
    ws2.cell(row=10, column=2, value="价格上限/下限").font = bold
    ws2.cell(
        row=10,
        column=3,
        value=f"{_fmt(curve.get('cap'))} / {_fmt(curve.get('floor'))}",
    )

    # ③ 中标与未中标
    ws3 = wb.create_sheet("中标清单")
    ws3["B2"] = "统一出清价结算（并列报价按申报比例分配并披露）"
    ws3["B2"].font = bold
    headers = ["主体", "资源类型", "申报容量（MW）", "中标容量（MW）", "报价（元/kW·年）"]
    for col, header in enumerate(headers, start=2):
        ws3.cell(row=4, column=col, value=header).font = bold
    for i, w in enumerate(formal.get("winners") or [], start=5):
        ws3.cell(row=i, column=2, value=str(w.get("subject_id")))
        ws3.cell(row=i, column=3, value=str(w.get("resource_type")))
        ws3.cell(row=i, column=4, value=_fmt(w.get("offered_capacity_mw")))
        ws3.cell(row=i, column=5, value=_fmt(w.get("cleared_mw")))
        ws3.cell(row=i, column=6, value=_fmt(w.get("offer_price")))
    ws4 = wb.create_sheet("未中标清单")
    ws4["B2"] = "未中标主体（含资格不合格）"
    ws4["B2"].font = bold
    headers = ["主体", "资源类型", "申报容量（MW）", "报价（元/kW·年）"]
    for col, header in enumerate(headers, start=2):
        ws4.cell(row=4, column=col, value=header).font = bold
    for i, w in enumerate(formal.get("losers") or [], start=5):
        ws4.cell(row=i, column=2, value=str(w.get("subject_id")))
        ws4.cell(row=i, column=3, value=str(w.get("resource_type")))
        ws4.cell(row=i, column=4, value=_fmt(w.get("offered_capacity_mw")))
        ws4.cell(row=i, column=5, value=_fmt(w.get("offer_price")))

    # ④ 市场力与缓解
    ws5 = wb.create_sheet("市场力与缓解")
    ws5["B2"] = "Q-M1：披露不阻断；缓解＝关键供应商报价替换重清一次（两版并列）"
    ws5["B2"].font = bold
    power = formal.get("market_power") or {}
    rows5 = [
        ("HHI", power.get("hhi")),
        ("最大主体占比", power.get("max_share")),
        ("关键供应商", "、".join(power.get("pivotal_subjects") or []) or "无"),
    ]
    for i, (label, value) in enumerate(rows5, start=4):
        ws5.cell(row=i, column=2, value=label).font = bold
        ws5.cell(
            row=i,
            column=3,
            value=(
                _fmt(value)
                if isinstance(value, (int, float)) and not isinstance(value, bool)
                else str(value)
            ),
        )
    mitigation = formal.get("mitigation")
    if mitigation:
        ws5.cell(row=8, column=2, value="缓解对比").font = bold
        for i, (label, value) in enumerate(
            [
                ("缓解前出清价", mitigation.get("price_before")),
                ("缓解后出清价", mitigation.get("price_after")),
                ("缓解参考成本", mitigation.get("reference_cost")),
                ("替换报价主体", "、".join(mitigation.get("replaced_subjects") or [])),
            ],
            start=9,
        ):
            ws5.cell(row=i, column=2, value=label)
            ws5.cell(
                row=i,
                column=3,
                value=(
                    _fmt(value)
                    if isinstance(value, (int, float))
                    and not isinstance(value, bool)
                    else str(value)
                ),
            )

    # ⑤ 披露
    ws6 = wb.create_sheet("披露")
    ws6["B2"] = "口径与边界（如实披露）"
    ws6["B2"].font = bold
    notes = [
        "本报告是已冻结出清结果（一次性不可变记录）的业务投影：不重算出清、不读台账最新版本。",
        "统一出清价＝边际主体报价；并列报价段按申报容量比例分配并披露（Q-M2）。",
        "触顶出清（供给在价格上限处仍不足最低采购要求）＝如实披露缺口，不静默成功（Q-M5）。",
        "市场力指标为披露性质，阈值判读随规则版本（Q-M1：HHI/最大占比/关键供应商）。",
        "需求量在曲线两端饱和（单区统一出清的诚实简化）；正式结果为缓解前口径。",
    ]
    for i, note in enumerate(notes, start=4):
        ws6.cell(row=i, column=2, value=note)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _html_bytes(
    *, formal: dict[str, Any], excel_name: str, excel_sha: str
) -> bytes:
    status = _STATUS_LABELS.get(str(formal.get("status")), formal.get("status"))
    kv = [
        ("出清状态", status),
        ("统一出清价", _fmt(formal.get("clearing_price")) + " 元/kW·年"),
        ("中标容量", _fmt(formal.get("procured_mw")) + " MW"),
        ("容量缺口", _fmt(formal.get("shortfall_mw")) + " MW"),
        ("目标容量", _fmt(formal.get("target_mw")) + " MW"),
        ("采购完成率", _fmt(formal.get("completion_ratio"))),
        ("规则版本", formal.get("rule_version")),
        ("确认人／冻结时间", f"{formal.get('confirmer')}／{formal.get('frozen_at')}"),
    ]
    winner_rows = "".join(
        f"<tr><td>{w.get('subject_id')}</td><td class='r'>{_fmt(w.get('offered_capacity_mw'))}</td>"
        f"<td class='r'>{_fmt(w.get('cleared_mw'))}</td><td class='r'>{_fmt(w.get('offer_price'))}</td></tr>"
        for w in formal.get("winners") or []
    )
    html = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<title>容量市场仿真出清报告</title></head>
<body>
<h1>容量市场仿真出清报告</h1>
<p>权威工件：{excel_name}（SHA-256：{excel_sha[:16]}…）；本页为业务阅读版，不替代工件。</p>
<h2>结论</h2>
<table border="1" cellpadding="4">{''.join(f'<tr><td>{k}</td><td>{v}</td></tr>' for k, v in kv)}</table>
<h2>中标清单（统一出清价）</h2>
<table border="1" cellpadding="4"><tr><th>主体</th><th>申报（MW）</th><th>中标（MW）</th><th>报价</th></tr>{winner_rows}</table>
<p>本报告为市场域业务投影：不修改机器结果，不重算出清；来源＝已冻结正式结果（一次性不可变）。</p>
</body></html>
"""
    return html.encode("utf-8")


def export_market_report(
    *, formal: dict[str, Any], output_dir: Path | str
) -> dict[str, Any]:
    """发布出清报告集（Excel＋HTML＋manifest 最后写；同身份重放按哈希）。
    只接受已冻结正式结果（formal.clearing_ref 在场）。"""
    output = Path(output_dir)
    clearing_ref = str(formal.get("clearing_ref") or "")
    if not clearing_ref:
        raise MarketReportError(
            "MARKET_REPORT_INPUT_INVALID", "出清记录缺 clearing_ref（未冻结）"
        )
    output.mkdir(parents=True, exist_ok=True)
    xlsx = _excel_bytes(formal=formal)
    xlsx_sha = _sha256(xlsx)
    html = _html_bytes(formal=formal, excel_name=_REPORT_XLSX, excel_sha=xlsx_sha)
    entries = [
        {
            "name": _REPORT_XLSX,
            "sha256": xlsx_sha,
            "size": len(xlsx),
            "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        },
        {
            "name": _REPORT_HTML,
            "sha256": _sha256(html),
            "size": len(html),
            "content_type": "text/html; charset=utf-8",
        },
    ]
    result_sha = _sha256(_canonical(formal))
    report_manifest = {
        "market_report_schema_version": "1.0",
        "generator": _GENERATOR,
        "clearing_ref": clearing_ref,
        "rule_version": formal.get("rule_version"),
        "result_sha256": result_sha,
        "files": entries,
        "exported_at": datetime.now(UTC).isoformat(),
    }
    manifest_bytes = _canonical(report_manifest)
    existing = output / _MANIFEST_FILE
    if existing.is_file():
        try:
            previous = json.loads(existing.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            raise MarketReportError(
                "MARKET_REPORT_EXISTS_DIFFERENT",
                f"既有市场报告清单不可读取：{exc}",
            ) from exc
        if previous.get("files") == report_manifest["files"] and previous.get(
            "result_sha256"
        ) == report_manifest["result_sha256"]:
            return {
                "replayed": True,
                "clearing_ref": clearing_ref,
                "output_dir": str(output),
                "files": [entry["name"] for entry in entries],
            }
        raise MarketReportError(
            "MARKET_REPORT_EXISTS_DIFFERENT",
            "同身份目录已存在不同的市场报告集；不覆盖既有报告",
        )
    for entry in entries:
        if (output / entry["name"]).exists():
            raise MarketReportError(
                "MARKET_REPORT_EXISTS_DIFFERENT", f"报告文件已存在：{entry['name']}"
            )
    from capacity_credit.application.atomic_io import atomic_write_bytes

    for entry in entries:
        name = entry["name"]
        data = xlsx if name == _REPORT_XLSX else html
        atomic_write_bytes(output / name, data, immutable=True)
    atomic_write_bytes(existing, manifest_bytes + b"\n", immutable=True)
    return {
        "replayed": False,
        "clearing_ref": clearing_ref,
        "output_dir": str(output),
        "files": [entry["name"] for entry in entries],
    }


__all__ = ["MarketReportError", "export_market_report"]
