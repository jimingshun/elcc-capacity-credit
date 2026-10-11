"""T24-A: table-first Excel renderer (openpyxl).

Consumes ONLY the shared ReportView and its display fields (labels,
units, display values and number formats originate from mapping_v1 via
report_model). Sheet names/order follow the product design 16.2 fixed
13-sheet list; non-applicable sheets are PRESERVED with a stable note.
Blocks are written through a section writer: every block returns the
next free row, blocks never overwrite each other. Security: strings
starting with = + - @ are prefixed with ' (strings only - real negative
numbers stay numeric cells); no macros, no formulas, no external links.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.chart import BarChart, LineChart, Reference  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.worksheet.properties import PageSetupProperties  # type: ignore[import-untyped]

from capacity_credit.reporting.report_model import ReportView

# Product design 16.2 fixed sheet list - strict order, all 13 present.
FIXED_SHEET_ORDER: tuple[str, ...] = (
    "说明与版本",
    "参数摘要",
    "数据质量",
    "系统充裕度",
    "ELCC",
    "EFC",
    "MRI",
    "储能时长",
    "渗透率",
    "敏感性",
    "风险时段",
    "搜索与收敛",
    "来源与审计",
)

_METHOD_SHEET: dict[str, str] = {
    "elcc": "ELCC",
    "efc": "EFC",
    "mri": "MRI",
    "rmri": "MRI",
}

_WARNING_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
_ERROR_FILL = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
_HEADER_FONT = Font(bold=True)
_BANNER_FONT = Font(bold=True, color="9C0006")


def _injection_safe(text: Any) -> Any:
    """Formula-injection guard: ONLY strings starting with = + - @ get a
    leading apostrophe; numbers (including real negatives) stay numeric."""
    if isinstance(text, str) and text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _title(sheet: Any, row: int, text: str) -> int:
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = _HEADER_FONT
    return row + 1


def _key_value_block(sheet: Any, start_row: int, pairs: list[tuple[str, Any]]) -> int:
    """Key/value block; returns the next free row (one blank separator
    after the block)."""
    row = start_row
    for key, value in pairs:
        sheet.cell(row=row, column=1, value=key)
        cell = sheet.cell(row=row, column=2, value=_injection_safe(value))
        cell.number_format = getattr(value, "number_format", None) or "General"
        row += 1
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 46
    return row + 1


def _table_block(
    sheet: Any,
    start_row: int,
    rows: list[dict[str, Any]],
    columns: list[str],
    *,
    widths: dict[str, float] | None = None,
) -> int:
    """Table block with header; returns the next free row (one blank
    separator after the block)."""
    if not rows:
        return start_row + 1
    for column, header in enumerate(columns, start=1):
        cell = sheet.cell(row=start_row, column=column, value=header)
        cell.font = _HEADER_FONT
    wrap = Alignment(wrap_text=True, vertical="top")
    for row_index, row in enumerate(rows, start=start_row + 1):
        for column, key in enumerate(columns, start=1):
            cell = sheet.cell(row=row_index, column=column, value=_injection_safe(row.get(key)))
            cell.alignment = wrap
    sheet.auto_filter.ref = (
        f"A{start_row}:{get_column_letter(len(columns))}{start_row + len(rows)}"
    )
    for column in range(1, len(columns) + 1):
        width = sheet.column_dimensions[get_column_letter(column)].width or 18
        sheet.column_dimensions[get_column_letter(column)].width = max(width, 22)
    if widths:
        for key, width in widths.items():
            if key in columns:
                column = columns.index(key) + 1
                sheet.column_dimensions[get_column_letter(column)].width = width
    return start_row + len(rows) + 2


_DISPLAY_COLUMN_WIDTHS: dict[str, float] = {
    "field_id": 30.0,
    "label": 26.0,
    "raw_value": 16.0,
    "display_unit": 14.0,
    "null_status": 18.0,
}


def _display_block(
    sheet: Any,
    start_row: int,
    fields: list[Any],
    columns: list[str],
) -> int:
    """Display-field table: raw values stay NUMERIC cells with the mapped
    number format; labels/units/null semantics come from the shared model.
    THE single shared layout entry: deterministic, bounded column widths
    and wrap text keep field_id / Chinese labels fully readable; no other
    block duplicates these style rules."""
    if not fields:
        return start_row + 1
    wrap = Alignment(wrap_text=True, vertical="top")
    for column, header in enumerate(columns, start=1):
        cell = sheet.cell(row=start_row, column=column, value=header)
        cell.font = _HEADER_FONT
    for column, key in enumerate(columns, start=1):
        width = _DISPLAY_COLUMN_WIDTHS.get(key)
        if width is not None:
            sheet.column_dimensions[get_column_letter(column)].width = width
    for row_index, field in enumerate(fields, start=start_row + 1):
        values = {
            "field_id": field.field_id,
            "label": field.label,
            "raw_value": field.raw_value,
            "display_unit": field.display_unit or "",
            "null_status": field.null_status,
        }
        for column, key in enumerate(columns, start=1):
            cell = sheet.cell(row=row_index, column=column, value=_injection_safe(values.get(key)))
            if key in ("field_id", "label"):
                cell.alignment = wrap
            if key == "raw_value" and field.number_format is not None:
                cell.number_format = field.number_format
    return start_row + len(fields) + 2


def _note_block(sheet: Any, start_row: int, note: str) -> int:
    sheet.cell(row=start_row, column=1, value=note)
    return start_row + 2


def _generator_version() -> str:
    """THE generator version from the report-model module (one literal,
    never a drifted copy inside the sheet)."""
    from capacity_credit.reporting.report_model import REPORT_GENERATOR_VERSION

    return REPORT_GENERATOR_VERSION


def render_excel(view: ReportView, path: str | None = None) -> bytes:
    """Render the 13 fixed sheets from the shared view; returns bytes
    (and writes them when a path is given - the unified export entry is
    the only sanctioned publisher)."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    capacity_fields = [f for f in view.display_fields]
    capacity_columns = ["field_id", "label", "raw_value", "display_unit", "null_status"]

    # print layout: wide tables print LANDSCAPE with fitToWidth=1
    # (fields and values stay on the same print page); screen column
    # widths are untouched; long text wraps instead of overflowing.
    _WIDE_SHEETS = {
        "数据质量",
        "系统充裕度",
        "ELCC",
        "EFC",
        "MRI",
        "风险时段",
        "搜索与收敛",
        "来源与审计",
    }
    for sheet_name in FIXED_SHEET_ORDER:
        sheet = workbook.create_sheet(title=sheet_name)
        if sheet_name in _WIDE_SHEETS:
            sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        row = 1
        if sheet_name == "说明与版本":
            if view.use_boundary_warning:
                cell = sheet.cell(row=row, column=1, value=view.use_boundary_warning)
                cell.font = _BANNER_FONT
                cell.fill = _ERROR_FILL if view.use_boundary_code == "diagnostic" else _WARNING_FILL
                row += 2
            row = _title(sheet, row, "说明与版本")
            row = _key_value_block(
                sheet,
                row,
                [
                    ("report_mapping_version", view.report_mapping_version),
                    ("run_id", view.run_id),
                    ("result_id", view.result_id),
                    ("result_status", view.result_status),
                    ("method", view.method),
                    ("method_profile", view.method_profile),
                    ("report_generator_version", _generator_version()),
                ],
            )
            # C-night/F: the deterministic business-reading notes ride on
            # the SAME published workbook (no second file, no second
            # generator); machine fields/units above are untouched
            if view.business_notes:
                row = _title(
                    sheet,
                    row,
                    "业务解读（"
                    + str(
                        view.business_notes.get("method_label")
                        or view.business_notes.get("method", "")
                    )
                    + "）",
                )
                for section in view.business_notes.get("sections", []):
                    row = _title(sheet, row, str(section.get("heading", "")))
                    for paragraph in section.get("paragraphs", []):
                        row = _note_block(sheet, row, str(paragraph))
                    table = section.get("table")
                    if table:
                        columns = [str(c) for c in table.get("columns", [])]
                        dict_rows = [
                            {columns[i]: str(cell) for i, cell in enumerate(table_row)}
                            for table_row in table.get("rows", [])
                        ]
                        row = _table_block(sheet, row, dict_rows, columns)
        elif sheet_name == "参数摘要":
            row = _title(sheet, row, "参数摘要")
            row = _key_value_block(
                sheet,
                row,
                [
                    ("resource_or_portfolio_id", view.resource_or_portfolio_id),
                    ("primary_risk_metric", view.primary_risk_metric),
                    ("primary_risk_metric_unit", view.primary_risk_metric_unit),
                    ("base_risk", view.base_risk),
                    ("candidate_risk", view.candidate_risk),
                    ("matched_risk", view.matched_risk),
                    ("risk_convergence", view.risk_convergence),
                    ("search_convergence", view.search_convergence),
                ],
            )
        elif sheet_name == "数据质量":
            row = _title(sheet, row, "数据质量")
            row = _key_value_block(
                sheet,
                row,
                [
                    ("effective_data_grade", view.effective_data_grade),
                    ("警告数", len(view.warnings)),
                    ("阻断性诊断数", len(view.blocking_diagnostics)),
                ],
            )
            row = _title(sheet, row, "空值语义与关键字段状态")
            row = _table_block(
                sheet,
                row,
                [
                    {
                        "field": entry.get("field_id", ""),
                        "label": entry.get("label", ""),
                        "value": entry.get("value"),
                        "null_status": entry.get("null_status", ""),
                    }
                    for entry in view.key_field_statuses
                ],
                ["field", "label", "value", "null_status"],
            )
            if view.warnings:
                row = _title(sheet, row, "正式警告")
                row = _table_block(
                    sheet,
                    row,
                    [
                        {"code": w.get("code", ""), "severity": w.get("severity", ""), "message": w.get("message", "")}
                        for w in view.warnings
                    ],
                    ["code", "severity", "message"],
                )
        elif sheet_name == "系统充裕度":
            row = _title(sheet, row, "系统充裕度")
            if view.adequacy is not None and view.adequacy.rows:
                row = _key_value_block(
                    sheet,
                    row,
                    [
                        ("总行数 total_rows", view.adequacy.total_rows),
                        ("已展示行数 displayed_rows", view.adequacy.displayed_rows),
                    ],
                )
                table_start = row
                row = _table_block(
                    sheet,
                    row,
                    list(view.adequacy.rows),
                    ["role", "metric", "value", "unit", "primary_metric", "result_status"],
                )
                # T24-C enhanced chart: primary-risk role comparison. The
                # chart references a dedicated chart-data zone written
                # from the SAME validated view rows (no second data
                # interpretation layer, no mixed-metric references).
                primary_rows = [
                    (r.get("role"), r.get("value"))
                    for r in view.adequacy.rows
                    if r.get("primary_metric") == "yes" and r.get("value") is not None
                ]
                if primary_rows:
                    sheet.cell(row=row, column=1, value="图表数据（主指标）")
                    chart_start = row
                    for index, (role, value) in enumerate(primary_rows):
                        sheet.cell(row=chart_start + 1 + index, column=1, value=role)
                        sheet.cell(
                            row=chart_start + 1 + index, column=2, value=value
                        )
                    chart = BarChart()
                    chart.type = "col"
                    chart.title = (
                        f"基准/候选/匹配 主风险对比（{view.primary_risk_metric}）"
                    )
                    chart.y_axis.title = view.primary_risk_metric_unit
                    chart.x_axis.title = "角色"
                    data_ref = Reference(
                        sheet,
                        min_col=2,
                        min_row=chart_start + 1,
                        max_row=chart_start + len(primary_rows),
                    )
                    cats_ref = Reference(
                        sheet,
                        min_col=1,
                        min_row=chart_start + 1,
                        max_row=chart_start + len(primary_rows),
                    )
                    chart.add_data(data_ref, titles_from_data=False)
                    chart.set_categories(cats_ref)
                    sheet.add_chart(chart, f"A{chart_start + len(primary_rows) + 2}")
                    row = chart_start + len(primary_rows) + 18
            else:
                row = _note_block(sheet, row, "不适用/未提供：本方法不生成系统充裕度长表")
        elif sheet_name in ("ELCC", "EFC", "MRI"):
            method_sheet = _METHOD_SHEET.get(view.method)
            row = _title(sheet, row, sheet_name)
            if method_sheet == sheet_name:
                applicable = {
                    "reference_capacity_mw",
                    "accredited_capacity_mw",
                    "capacity_credit_fraction",
                    "capacity_credit_unit",
                    "result_applicability",
                    "capacity_value_type",
                }
                block_start = row
                row = _display_block(
                    sheet,
                    row,
                    [f for f in capacity_fields if f.field_id in applicable],
                    capacity_columns,
                )
                # T24-C enhanced chart: capacity value (reference vs
                # accredited) from the SAME validated display block
                ref_row = next(
                    (
                        block_start + index + 1
                        for index, f in enumerate(capacity_fields)
                        if f.field_id == "reference_capacity_mw" and f.field_id in applicable
                    ),
                    None,
                )
                acc_row = next(
                    (
                        block_start + index + 1
                        for index, f in enumerate(capacity_fields)
                        if f.field_id == "accredited_capacity_mw" and f.field_id in applicable
                    ),
                    None,
                )
                if ref_row is not None and acc_row is not None:
                    chart = BarChart()
                    chart.type = "col"
                    chart.title = "容量价值（认证容量 vs 参考容量）"
                    chart.y_axis.title = "MW"
                    chart.x_axis.title = "容量项"
                    data_ref = Reference(sheet, min_col=3, min_row=ref_row, max_row=acc_row)
                    cats_ref = Reference(sheet, min_col=1, min_row=ref_row, max_row=acc_row)
                    chart.add_data(data_ref, titles_from_data=False)
                    chart.set_categories(cats_ref)
                    sheet.add_chart(chart, f"A{row + 1}")
                    row += 16
                if view.method in ("mri", "rmri") and sheet_name == "MRI":
                    row = _display_block(
                        sheet,
                        row,
                        [
                            f
                            for f in capacity_fields
                            if f.field_id in ("rmri_status", "rmri_value", "mri_value")
                        ],
                        capacity_columns,
                    )
            else:
                row = _note_block(sheet, row, f"不适用/未提供：方法 {view.method} 不生成此工作表")
        elif sheet_name in ("储能时长", "渗透率", "敏感性"):
            row = _title(sheet, row, sheet_name)
            row = _note_block(sheet, row, "不适用/未提供：T24-A 当前方法集不生成此工作表（机器结果结构保留）")
        elif sheet_name == "风险时段":
            row = _title(sheet, row, "风险时段")
            if view.risk_events is not None:
                row = _key_value_block(
                    sheet,
                    row,
                    [
                        ("总事件行数 total_rows", view.risk_events.total_rows),
                        ("已展示行数 displayed_rows", view.risk_events.displayed_rows),
                        ("权威源文件", view.risk_events.source),
                    ],
                )
                row = _table_block(
                    sheet,
                    row,
                    list(view.risk_events.rows),
                    [
                        "event_time", "location", "deficiency_power_mw",
                        "deficiency_energy_mwh", "scenario_id", "sample_index",
                        "duration_hours",
                    ],
                )
            else:
                row = _note_block(sheet, row, "不适用/未提供：无风险事件")
        elif sheet_name == "搜索与收敛":
            row = _title(sheet, row, "搜索与收敛")
            if view.search_trace is not None and view.search_trace.rows:
                table_start = row
                row = _table_block(
                    sheet,
                    row,
                    list(view.search_trace.rows),
                    [
                        "call_id", "iteration", "action", "capacity_mw", "reference_risk",
                        "trial_risk", "difference_mean", "difference_se", "sample_count",
                        "bracket_low_after", "bracket_high_after", "risk_convergence",
                        "estimate_mw", "trace_status",
                    ],
                )
                # T24-C enhanced chart: search capacity path by call_id,
                # referencing the SAME validated table cells
                point_rows = [
                    (table_start + index + 1, r.get("call_id"), r.get("capacity_mw"))
                    for index, r in enumerate(view.search_trace.rows)
                    if r.get("call_id") is not None and r.get("capacity_mw") is not None
                ]
                if len(point_rows) >= 2:
                    chart = LineChart()
                    chart.title = "搜索轨迹：call_id 上的容量试探（capacity_mw）"
                    chart.y_axis.title = "capacity_mw"
                    chart.x_axis.title = "call_id"
                    data_ref = Reference(
                        sheet,
                        min_col=4,
                        min_row=point_rows[0][0],
                        max_row=point_rows[-1][0],
                    )
                    cats_ref = Reference(
                        sheet,
                        min_col=1,
                        min_row=point_rows[0][0],
                        max_row=point_rows[-1][0],
                    )
                    chart.add_data(data_ref, titles_from_data=False)
                    chart.set_categories(cats_ref)
                    sheet.add_chart(chart, f"A{row + 1}")
                    row += 16
            else:
                row = _note_block(
                    sheet, row, view.search_trace.note if view.search_trace else "不适用/未提供：无搜索轨迹"
                )
        elif sheet_name == "来源与审计":
            row = _title(sheet, row, "来源与审计")
            row = _key_value_block(
                sheet,
                row,
                [
                    ("rerun_capability", view.reproducibility.get("rerun_capability", "")),
                    (
                        "missing_reproducibility_inputs",
                        ", ".join(view.reproducibility.get("missing_reproducibility_inputs") or []),
                    ),
                ],
            )
            row = _title(sheet, row, "输入产物（七项）")
            row = _table_block(
                sheet,
                row,
                [
                    {"artifact": record.name, "sha256": record.sha256, "size": record.size}
                    for record in sorted(view.source_files.values(), key=lambda r: r.name)
                ],
                ["artifact", "sha256", "size"],
                widths={"sha256": 74.0},
            )

    buffer = io.BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    if path is not None:
        Path(path).write_bytes(data)
    return data


__all__ = ["FIXED_SHEET_ORDER", "render_excel"]
